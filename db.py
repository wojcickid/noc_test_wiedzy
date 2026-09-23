"""Wspólny dostęp do bazy SQLite używany przez aplikację i skrypty pomocnicze."""

import csv
import io
import json
import os
import re
import secrets
import sqlite3
import string
from contextlib import contextmanager
from datetime import datetime, timedelta

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PLIK = os.path.join(BASE_DIR, "baza.db")

# Bez znaków łatwych do pomylenia przy przepisywaniu (0/O, 1/I/L) — używane
# zarówno do tokenów, jak i do generowanego hasła panelu administracyjnego.
ALFABET_KODOW = "".join(sorted(set(string.ascii_uppercase + string.digits) - set("0O1IL")))

WYMAGANE_KOLUMNY_PYTAN = ["tresc_pytania", "opcja_a", "opcja_b", "opcja_c", "opcja_d", "odpowiedz"]
# Opcjonalna — treść pokazywana przy wyniku uczestnika, wyjaśniająca poprawną
# odpowiedź (E4). Plik bez tej kolumny importuje się normalnie.
KOLUMNA_WYJASNIENIE = "wyjasnienie"
ODPOWIEDZI_DOZWOLONE = {"A", "B", "C", "D"}

# Zapas czasu po upłynięciu limitu (F1) — bufor na opóźnienie sieci przy
# auto-wysyłce formularza dokładnie w momencie 0:00 na liczniku.
TOLERANCJA_SEKUNDY = 5


class BladImportu(Exception):
    pass

SCHEMAT_TABELE = """
CREATE TABLE IF NOT EXISTS testy (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    nazwa TEXT NOT NULL UNIQUE,
    data_importu TEXT NOT NULL,
    liczba_pytan_do_losowania INTEGER NOT NULL DEFAULT 20,
    limit_czasu_min INTEGER,
    tryb_szczegolow TEXT NOT NULL DEFAULT 'po_zamknieciu',
    szczegoly_od TEXT,
    wynik_widoczny TEXT NOT NULL DEFAULT 'od_razu',
    dostepny_od TEXT,
    dostepny_do TEXT,
    zamkniety INTEGER NOT NULL DEFAULT 0,
    prog_zaliczenia INTEGER DEFAULT 80,
    mail_temat TEXT,
    mail_tresc TEXT,
    przypomnienie_temat TEXT,
    przypomnienie_tresc TEXT
);

CREATE TABLE IF NOT EXISTS pytania (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    test_id INTEGER NOT NULL REFERENCES testy(id),
    tresc_pytania TEXT NOT NULL,
    opcja_a TEXT NOT NULL,
    opcja_b TEXT NOT NULL,
    opcja_c TEXT NOT NULL,
    opcja_d TEXT NOT NULL,
    odpowiedz TEXT NOT NULL,
    wyjasnienie TEXT
);

CREATE TABLE IF NOT EXISTS tokeny (
    token TEXT PRIMARY KEY,
    test_id INTEGER NOT NULL REFERENCES testy(id),
    wykorzystany INTEGER NOT NULL DEFAULT 0,
    data_utworzenia TEXT NOT NULL,
    data_wykorzystania TEXT,
    imie TEXT,
    email TEXT,
    mail_wyslany_at TEXT,
    przypomnienie_wyslane_at TEXT,
    mail_status TEXT,
    mail_blad TEXT
);

-- Ustawienia całej instalacji edytowane w panelu (Etap 6): nazwa nadawcy
-- maili i adres aplikacji do linków. Dane logowania SMTP są celowo poza bazą
-- (config.json), żeby hasło nie trafiło do kopii bazy pobieranej z panelu.
CREATE TABLE IF NOT EXISTS ustawienia_aplikacji (
    klucz TEXT PRIMARY KEY,
    wartosc TEXT
);

CREATE TABLE IF NOT EXISTS odpowiedzi_uzytkownika (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    token TEXT NOT NULL REFERENCES tokeny(token),
    pytanie_id INTEGER NOT NULL REFERENCES pytania(id),
    odpowiedz TEXT NOT NULL,
    data_wyslania TEXT NOT NULL
);

-- Postęp podejścia do testu (Etap 2, L1) — jedna aktywna/zakończona próba na
-- token. Każda odpowiedź trafia od razu do odpowiedzi_uzytkownika, więc utrata
-- ciasteczka/sesji przeglądarki nie kasuje postępu (B1, B13) i test można
-- wznowić tym samym tokenem na dowolnym urządzeniu (E13).
CREATE TABLE IF NOT EXISTS podejscia (
    token TEXT PRIMARY KEY REFERENCES tokeny(token),
    test_id INTEGER NOT NULL REFERENCES testy(id),
    wybrane_pytania TEXT NOT NULL,
    liczba_pytan INTEGER NOT NULL,
    indeks_pytania INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'w_trakcie',
    data_rozpoczecia TEXT NOT NULL,
    data_zakonczenia TEXT
);

CREATE INDEX IF NOT EXISTS idx_pytania_test ON pytania(test_id);
CREATE INDEX IF NOT EXISTS idx_tokeny_test ON tokeny(test_id);
CREATE INDEX IF NOT EXISTS idx_odpowiedzi_token ON odpowiedzi_uzytkownika(token);
CREATE INDEX IF NOT EXISTS idx_podejscia_test ON podejscia(test_id);
-- Jeden wiersz na parę (token, pytanie) — gwarantuje w SQLite, że dwa
-- równoczesne zapisy tej samej odpowiedzi się nie zduplikują (B8).
CREATE UNIQUE INDEX IF NOT EXISTS idx_odpowiedzi_unikalne ON odpowiedzi_uzytkownika(token, pytanie_id);
"""

# Widoki są odtwarzane przy każdym starcie (DROP + CREATE, nie IF NOT EXISTS),
# żeby zmiany w ich definicji zawsze się zastosowały — nie trzymają danych,
# więc nie ma ryzyka utraty czegokolwiek.
SCHEMAT_WIDOKI = """
DROP VIEW IF EXISTS arkusz_wynikow;
-- Odpowiednik dawnego "arkusza wyników" (wyniki_<token>.xlsx) — pytanie po
-- pytaniu, z odpowiedzią użytkownika i poprawną odpowiedzią (litera + treść).
-- Tylko podejścia zakończone (L1) — w trakcie trwania testu nie pokazujemy
-- cząstkowych wyników nikomu, kto poda/sprawdzi ten sam token po drodze.
CREATE VIEW arkusz_wynikow AS
SELECT
    ou.token AS token,
    p.tresc_pytania AS tresc_pytania,
    ou.odpowiedz AS odpowiedz_uzytkownika,
    CASE ou.odpowiedz
        WHEN 'A' THEN p.opcja_a WHEN 'B' THEN p.opcja_b
        WHEN 'C' THEN p.opcja_c WHEN 'D' THEN p.opcja_d
    END AS tresc_odpowiedzi_uzytkownika,
    p.odpowiedz AS poprawna_odpowiedz,
    CASE p.odpowiedz
        WHEN 'A' THEN p.opcja_a WHEN 'B' THEN p.opcja_b
        WHEN 'C' THEN p.opcja_c WHEN 'D' THEN p.opcja_d
    END AS tresc_poprawnej_odpowiedzi,
    p.wyjasnienie AS wyjasnienie,
    ou.data_wyslania AS data_wyslania
FROM odpowiedzi_uzytkownika ou
JOIN pytania p ON p.id = ou.pytanie_id
JOIN podejscia pj ON pj.token = ou.token
WHERE pj.status IN ('zakonczone', 'czas_minal');

DROP VIEW IF EXISTS zbiorcze_wyniki;
-- Odpowiednik dawnego wyniki_testu.xlsx — jeden wiersz na zakończone podejście.
-- Mianownik (wszystkie) to liczba wylosowanych pytań z podejscia, nie COUNT()
-- udzielonych odpowiedzi (L3) — te dwie liczby są równe dla normalnie
-- ukończonego podejścia, ale tylko ta pierwsza jest poprawna, gdyby kiedyś
-- się rozjechały (np. ręczna ingerencja w dane).
CREATE VIEW zbiorcze_wyniki AS
SELECT
    pj.token AS token,
    tk.imie AS imie,
    tk.email AS email,
    t.id AS test_id,
    t.nazwa AS test,
    pj.liczba_pytan AS wszystkie,
    (
        SELECT COUNT(*) FROM odpowiedzi_uzytkownika ou
        JOIN pytania p ON p.id = ou.pytanie_id
        WHERE ou.token = pj.token AND ou.odpowiedz = p.odpowiedz
    ) AS poprawne,
    ROUND(
        100.0 * (
            SELECT COUNT(*) FROM odpowiedzi_uzytkownika ou
            JOIN pytania p ON p.id = ou.pytanie_id
            WHERE ou.token = pj.token AND ou.odpowiedz = p.odpowiedz
        ) / pj.liczba_pytan, 1
    ) AS procent,
    pj.data_zakonczenia AS data_wyslania,
    pj.data_rozpoczecia AS data_rozpoczecia,
    pj.status AS status
FROM podejscia pj
JOIN tokeny tk ON tk.token = pj.token
JOIN testy t ON t.id = tk.test_id
WHERE pj.status IN ('zakonczone', 'czas_minal');
"""


def _migruj_tabele(conn):
    """Dodaje kolumny do już istniejących baz, jeśli schemat ewoluował od czasu
    ich utworzenia — tabele (w przeciwieństwie do widoków) trzymają dane, więc
    nie można ich po prostu odtworzyć od nowa."""
    # Etap 6 (L6) — osobne kolumny imie/email zamiast jednego pola `przypisany`
    # oraz stan wysyłki maili. Stara kolumna `przypisany` zostaje w starych
    # bazach (SQLite nie usuwa kolumn bez przebudowy tabeli), ale nie jest już
    # używana — jej zawartość jest jednorazowo rozdzielana na imie/email.
    kolumny_tokeny = {w["name"] for w in conn.execute("PRAGMA table_info(tokeny)").fetchall()}
    if "imie" not in kolumny_tokeny:
        conn.execute("ALTER TABLE tokeny ADD COLUMN imie TEXT")
        conn.execute("ALTER TABLE tokeny ADD COLUMN email TEXT")
        if "przypisany" in kolumny_tokeny:
            for wiersz in conn.execute(
                "SELECT token, przypisany FROM tokeny WHERE przypisany IS NOT NULL AND TRIM(przypisany) != ''"
            ).fetchall():
                imie, email = _rozdziel_przypisanie(wiersz["przypisany"])
                conn.execute("UPDATE tokeny SET imie = ?, email = ? WHERE token = ?", (imie, email, wiersz["token"]))
    for kolumna in ("mail_wyslany_at", "przypomnienie_wyslane_at", "mail_status", "mail_blad"):
        if kolumna not in kolumny_tokeny:
            conn.execute(f"ALTER TABLE tokeny ADD COLUMN {kolumna} TEXT")

    kolumny_testy = {w["name"] for w in conn.execute("PRAGMA table_info(testy)").fetchall()}
    if "liczba_pytan_do_losowania" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN liczba_pytan_do_losowania INTEGER NOT NULL DEFAULT 20")

    kolumny_pytania = {w["name"] for w in conn.execute("PRAGMA table_info(pytania)").fetchall()}
    if "wyjasnienie" not in kolumny_pytania:
        conn.execute("ALTER TABLE pytania ADD COLUMN wyjasnienie TEXT")

    # Etap 4 (F1, F2) — ustawienia kontroli nad testem: limit czasu, widoczność
    # wyników/szczegółów, okno dostępności, zamknięcie testu.
    kolumny_testy = {w["name"] for w in conn.execute("PRAGMA table_info(testy)").fetchall()}
    if "limit_czasu_min" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN limit_czasu_min INTEGER")
    if "tryb_szczegolow" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN tryb_szczegolow TEXT NOT NULL DEFAULT 'po_zamknieciu'")
    if "szczegoly_od" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN szczegoly_od TEXT")
    if "wynik_widoczny" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN wynik_widoczny TEXT NOT NULL DEFAULT 'od_razu'")
    if "dostepny_od" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN dostepny_od TEXT")
    if "dostepny_do" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN dostepny_do TEXT")
    if "zamkniety" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN zamkniety INTEGER NOT NULL DEFAULT 0")
    # Etap 5 (E2) — próg zaliczenia w procentach, NULL = bez progu. Domyślne
    # 80% dostają też testy istniejące przed migracją (decyzja usera z Etapu 5).
    if "prog_zaliczenia" not in kolumny_testy:
        conn.execute("ALTER TABLE testy ADD COLUMN prog_zaliczenia INTEGER DEFAULT 80")
    # Etap 6 (F3, E6) — własne szablony maili per test, NULL = szablon domyślny.
    for kolumna in ("mail_temat", "mail_tresc", "przypomnienie_temat", "przypomnienie_tresc"):
        if kolumna not in kolumny_testy:
            conn.execute(f"ALTER TABLE testy ADD COLUMN {kolumna} TEXT")

    _odtworz_historyczne_podejscia(conn)


def _odtworz_historyczne_podejscia(conn):
    """Etap 2 (L1) wprowadził tabelę `podejscia` — starsze odpowiedzi zapisane
    jeszcze w modelu "wszystko naraz w /wynik" nie mają dla siebie wiersza w tej
    tabeli. Bez tego zniknęłyby z widoków wyników (które teraz wymagają
    podejscia.status='zakonczone'), więc odtwarzamy brakujące wiersze jako
    zakończone podejścia na podstawie istniejących odpowiedzi. Idempotentne —
    po pierwszym uruchomieniu każdy taki token ma już swój wiersz."""
    brakujace_podejscia = conn.execute(
        """
        SELECT DISTINCT ou.token AS token, tk.test_id AS test_id
        FROM odpowiedzi_uzytkownika ou
        JOIN tokeny tk ON tk.token = ou.token
        LEFT JOIN podejscia pj ON pj.token = ou.token
        WHERE pj.token IS NULL
        """
    ).fetchall()
    for wiersz in brakujace_podejscia:
        pytania_id = [
            w["pytanie_id"]
            for w in conn.execute(
                "SELECT pytanie_id FROM odpowiedzi_uzytkownika WHERE token = ? ORDER BY id",
                (wiersz["token"],),
            ).fetchall()
        ]
        czasy = conn.execute(
            "SELECT MIN(data_wyslania) AS start, MAX(data_wyslania) AS koniec "
            "FROM odpowiedzi_uzytkownika WHERE token = ?",
            (wiersz["token"],),
        ).fetchone()
        conn.execute(
            """
            INSERT INTO podejscia
                (token, test_id, wybrane_pytania, liczba_pytan, indeks_pytania, status, data_rozpoczecia, data_zakonczenia)
            VALUES (?, ?, ?, ?, ?, 'zakonczone', ?, ?)
            """,
            (
                wiersz["token"],
                wiersz["test_id"],
                json.dumps(pytania_id),
                len(pytania_id),
                len(pytania_id),
                czasy["start"],
                czasy["koniec"],
            ),
        )


def polacz():
    conn = sqlite3.connect(DB_PLIK, timeout=10)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    conn.row_factory = sqlite3.Row
    return conn


def inicjalizuj():
    with polacz() as conn:
        conn.executescript(SCHEMAT_TABELE)
        _migruj_tabele(conn)
        conn.executescript(SCHEMAT_WIDOKI)
        conn.commit()
    conn.close()


@contextmanager
def baza():
    """Kontekst: otwiera połączenie, commituje po udanym bloku, zawsze zamyka."""
    conn = polacz()
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def losowy_kod(dlugosc=8):
    return "".join(secrets.choice(ALFABET_KODOW) for _ in range(dlugosc))


def _komorka_na_tekst(wartosc):
    """Zamienia wartość komórki openpyxl na tekst: pusta komórka (`None`) daje
    `""` zamiast napisu `"nan"`, a liczba całkowita zapisana jako float (np.
    `5.0`, bo kolumna miała gdzieś pustą komórkę) traci zbędne `.0` (B5)."""
    if wartosc is None:
        return ""
    if isinstance(wartosc, float) and wartosc.is_integer():
        return str(int(wartosc))
    return str(wartosc).strip()


def wczytaj_i_zwaliduj_plik_pytan(dane_pliku):
    """Parsuje plik .xlsx z pytaniami i waliduje go wiersz po wierszu (B6).
    Nagłówki są normalizowane (`strip().lower()`), więc `Tresc_pytania` albo
    spacja na końcu nazwy kolumny nie psują importu. Puste wiersze (np. odstęp
    na końcu arkusza) są pomijane bez błędu. Zwraca `(pytania, bledy)` —
    `pytania` to lista słowników gotowych do zapisu (do podglądu przed
    zatwierdzeniem, L5), `bledy` to lista czytelnych komunikatów, po jednym na
    problem. Niepusta lista `bledy` oznacza, że pliku nie da się zaimportować
    w całości — cały plik jest wtedy odrzucany (decyzja z Etapu 3)."""
    try:
        skoroszyt = openpyxl.load_workbook(filename=io.BytesIO(dane_pliku), read_only=True, data_only=True)
    except Exception:
        return [], ["Nie udało się odczytać pliku — sprawdź, czy to poprawny plik .xlsx."]

    arkusz = skoroszyt.active
    wiersze = arkusz.iter_rows(values_only=True)
    try:
        naglowki_surowe = next(wiersze)
    except StopIteration:
        return [], ["Plik jest pusty."]

    naglowki = [(str(h).strip().lower() if h is not None else "") for h in naglowki_surowe]
    brakujace = [k for k in WYMAGANE_KOLUMNY_PYTAN if k not in naglowki]
    if brakujace:
        return [], [f"W pliku brakuje kolumn: {', '.join(brakujace)}"]

    indeksy = {nazwa: naglowki.index(nazwa) for nazwa in WYMAGANE_KOLUMNY_PYTAN}
    indeks_wyjasnienia = naglowki.index(KOLUMNA_WYJASNIENIE) if KOLUMNA_WYJASNIENIE in naglowki else None

    pytania = []
    bledy = []
    for numer_wiersza, wiersz in enumerate(wiersze, start=2):
        if wiersz is None or all(komorka is None for komorka in wiersz):
            continue

        wartosci = {
            nazwa: (_komorka_na_tekst(wiersz[idx]) if idx < len(wiersz) else "")
            for nazwa, idx in indeksy.items()
        }
        for nazwa in WYMAGANE_KOLUMNY_PYTAN:
            if not wartosci[nazwa]:
                bledy.append(f"wiersz {numer_wiersza}: puste pole '{nazwa}'")

        odpowiedz = wartosci["odpowiedz"].strip().upper()
        if wartosci["odpowiedz"] and odpowiedz not in ODPOWIEDZI_DOZWOLONE:
            bledy.append(
                f"wiersz {numer_wiersza}: nieprawidłowa odpowiedź '{wartosci['odpowiedz']}' (dozwolone: A, B, C, D)"
            )

        wyjasnienie = None
        if indeks_wyjasnienia is not None and indeks_wyjasnienia < len(wiersz):
            wyjasnienie = _komorka_na_tekst(wiersz[indeks_wyjasnienia]) or None

        pytania.append(
            {
                "tresc_pytania": wartosci["tresc_pytania"],
                "opcja_a": wartosci["opcja_a"],
                "opcja_b": wartosci["opcja_b"],
                "opcja_c": wartosci["opcja_c"],
                "opcja_d": wartosci["opcja_d"],
                "odpowiedz": odpowiedz,
                "wyjasnienie": wyjasnienie,
            }
        )

    if not pytania and not bledy:
        bledy.append("Plik nie zawiera żadnych pytań.")

    return pytania, bledy


def importuj_pytania(nazwa_testu, pytania, liczba_pytan=20):
    """Tworzy nowy test i wczytuje do niego już zwalidowane pytania (patrz
    `wczytaj_i_zwaliduj_plik_pytan`). `liczba_pytan` to ile pytań losować przy
    każdym podejściu — przycinane do liczby pytań w pliku. Zwraca
    (test_id, liczba_pytan_w_pliku, faktyczna_liczba_do_losowania)."""
    if not pytania:
        raise BladImportu("Plik nie zawiera żadnych pytań.")

    liczba_pytan = max(1, min(int(liczba_pytan), len(pytania)))

    with baza() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO testy (nazwa, data_importu, liczba_pytan_do_losowania) VALUES (?, ?, ?)",
                (nazwa_testu, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), liczba_pytan),
            )
        except sqlite3.IntegrityError:
            raise BladImportu(f"Test o nazwie '{nazwa_testu}' już istnieje. Wybierz inną nazwę.")
        test_id = cur.lastrowid

        conn.executemany(
            """
            INSERT INTO pytania (test_id, tresc_pytania, opcja_a, opcja_b, opcja_c, opcja_d, odpowiedz, wyjasnienie)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    test_id,
                    p["tresc_pytania"], p["opcja_a"], p["opcja_b"], p["opcja_c"], p["opcja_d"],
                    p["odpowiedz"], p.get("wyjasnienie"),
                )
                for p in pytania
            ],
        )

    return test_id, len(pytania), liczba_pytan


# Celowo prosta walidacja adresu (L6): wyłapuje literówki typu brak @, spacja
# czy przecinek w adresie. Czy skrzynka istnieje, wie dopiero serwer odbiorcy
# — taki błąd wraca później jako zwrotka na skrzynkę nadawcy, nie do aplikacji.
WZORZEC_EMAIL = re.compile(r"^[^@\s;,<>()\"']+@[^@\s;,<>()\"']+\.[^@\s;,<>()\"'.][^@\s;,<>()\"']*$")
MAKS_DLUGOSC_IMIENIA = 200
MAKS_DLUGOSC_EMAILA = 254
MAKS_UCZESTNIKOW = 500

NAGLOWKI_IMIENIA = {"imie", "imię", "imie i nazwisko", "imię i nazwisko", "uczestnik", "osoba", "nazwa"}
NAGLOWKI_EMAILA = {"email", "e-mail", "mail", "adres e-mail", "adres email"}


def waliduj_uczestnika(imie, email):
    """Czyści i sprawdza parę (imię, e-mail) — zwraca ją z None zamiast
    pustych napisów albo rzuca ValueError z czytelnym komunikatem."""
    imie = " ".join((imie or "").split()) or None
    email = (email or "").strip() or None
    if imie is None and email is None:
        raise ValueError("brak imienia i adresu e-mail")
    if imie and len(imie) > MAKS_DLUGOSC_IMIENIA:
        raise ValueError(f"imię dłuższe niż {MAKS_DLUGOSC_IMIENIA} znaków")
    if email and (len(email) > MAKS_DLUGOSC_EMAILA or not WZORZEC_EMAIL.match(email)):
        raise ValueError(f"nieprawidłowy adres e-mail '{email}'")
    return imie, email


def parsuj_uczestnika(linia):
    """Jedna linia listy uczestników (L6) -> (imie, email). Przyjmuje
    `Imię Nazwisko;email`, a także tabulator/przecinek jako separator, sam
    e-mail, samo imię albo `Imię Nazwisko <email>` (format z Outlooka) —
    adresem jest to słowo, które zawiera @. Rzuca ValueError przy błędzie."""
    slowa = [s for s in re.split(r"[;,\t\s]+", linia.strip()) if s]
    adresy = [s.strip("<>()\"'") for s in slowa if "@" in s]
    imie = " ".join(s for s in slowa if "@" not in s)
    if len(adresy) > 1:
        raise ValueError("więcej niż jeden adres e-mail w linii")
    return waliduj_uczestnika(imie, adresy[0] if adresy else None)


def _rozdziel_przypisanie(tekst):
    """Migracja starego pola `przypisany` (L6) — jak parsuj_uczestnika, ale
    nigdy nie odrzuca wpisu: jeśli nie da się wyciągnąć poprawnego adresu,
    całość trafia do imienia, żeby żadne dane nie zginęły."""
    try:
        return parsuj_uczestnika(tekst)
    except ValueError:
        return " ".join(tekst.split())[:MAKS_DLUGOSC_IMIENIA] or None, None


def _sprawdz_liste_uczestnikow(uczestnicy_z_numerami, jednostka):
    """Wspólna część parsowania listy z pola tekstowego i z pliku: limit
    liczby osób i powtórzone adresy (dwa tokeny na jedną skrzynkę to prawie
    zawsze pomyłka przy wklejaniu). `jednostka` to „linia” albo „wiersz”."""
    bledy = []
    if len(uczestnicy_z_numerami) > MAKS_UCZESTNIKOW:
        bledy.append(f"Lista może zawierać maksymalnie {MAKS_UCZESTNIKOW} osób naraz.")
    widziane = {}
    for numer, (_, email) in uczestnicy_z_numerami:
        if email:
            klucz = email.lower()
            if klucz in widziane:
                bledy.append(f"{jednostka} {numer}: adres {email} powtarza się (pierwszy raz: {jednostka} {widziane[klucz]})")
            else:
                widziane[klucz] = numer
    return [u for _, u in uczestnicy_z_numerami], bledy


def parsuj_liste_uczestnikow(tekst):
    """Lista wklejona w panelu (jedna osoba na linię) -> (uczestnicy, bledy).
    Tak jak przy imporcie pytań, jeden błąd odrzuca całą listę — dzięki temu
    nie powstaje połowa tokenów, której potem trzeba szukać."""
    wynik, bledy = [], []
    for numer, linia in enumerate(tekst.splitlines(), start=1):
        if not linia.strip():
            continue
        try:
            wynik.append((numer, parsuj_uczestnika(linia)))
        except ValueError as e:
            bledy.append(f"linia {numer}: {e}")
    uczestnicy, bledy_listy = _sprawdz_liste_uczestnikow(wynik, "linia")
    return uczestnicy, bledy + bledy_listy


def _wiersze_z_csv(dane_pliku):
    try:
        tekst = dane_pliku.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel w polskim Windowsie zapisuje „CSV (rozdzielany przecinkami)” w cp1250.
        tekst = dane_pliku.decode("cp1250", errors="replace")
    try:
        separator = csv.Sniffer().sniff(tekst[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        separator = ";"
    return list(csv.reader(io.StringIO(tekst), delimiter=separator))


def wczytaj_plik_uczestnikow(dane_pliku, nazwa_pliku):
    """Lista uczestników z pliku .xlsx albo .csv (L6) -> (uczestnicy, bledy).
    Pierwszy wiersz to nagłówki — kolumna z imieniem (`imie`, `imię`, `imię i
    nazwisko`…) i/lub z adresem (`email`, `e-mail`, `mail`…), wielkość liter
    bez znaczenia; pozostałe kolumny są ignorowane."""
    if nazwa_pliku.lower().endswith(".csv"):
        wiersze = _wiersze_z_csv(dane_pliku)
    else:
        try:
            skoroszyt = openpyxl.load_workbook(filename=io.BytesIO(dane_pliku), read_only=True, data_only=True)
        except Exception:
            return [], ["Nie udało się odczytać pliku — sprawdź, czy to poprawny plik .xlsx albo .csv."]
        wiersze = [[_komorka_na_tekst(k) for k in w] for w in skoroszyt.active.iter_rows(values_only=True)]
    if not wiersze:
        return [], ["Plik jest pusty."]

    naglowki = [str(h or "").strip().lower() for h in wiersze[0]]
    idx_imie = next((i for i, h in enumerate(naglowki) if h in NAGLOWKI_IMIENIA), None)
    idx_email = next((i for i, h in enumerate(naglowki) if h in NAGLOWKI_EMAILA), None)
    if idx_imie is None and idx_email is None:
        return [], ["W pierwszym wierszu pliku brakuje nagłówków kolumn: 'imie' i/lub 'email'."]

    def pole(wiersz, idx):
        return str(wiersz[idx] or "") if idx is not None and idx < len(wiersz) else ""

    wynik, bledy = [], []
    for numer, wiersz in enumerate(wiersze[1:], start=2):
        imie, email = pole(wiersz, idx_imie), pole(wiersz, idx_email)
        if not imie.strip() and not email.strip():
            continue
        try:
            wynik.append((numer, waliduj_uczestnika(imie, email)))
        except ValueError as e:
            bledy.append(f"wiersz {numer}: {e}")
    uczestnicy, bledy_listy = _sprawdz_liste_uczestnikow(wynik, "wiersz")
    if not uczestnicy and not bledy:
        bledy.append("Plik nie zawiera żadnych uczestników.")
    return uczestnicy, bledy + bledy_listy


def opis_uczestnika(imie, email):
    """Jeden napis do wyświetlania: `Imię <email>`, samo imię albo sam e-mail."""
    if imie and email:
        return f"{imie} <{email}>"
    return imie or email or ""


def wygeneruj_tokeny(test_id, liczba, dlugosc=8, uczestnicy=None):
    """Generuje `liczba` nowych, unikalnych tokenów dla wskazanego testu.
    Jeśli podano `uczestnicy` (lista par (imie, email) o długości `liczba`),
    i-ty token dostaje i-tą osobę. Zwraca listę {"token", "imie", "email"}."""
    if dlugosc < 4 or dlugosc > 20:
        raise ValueError("Długość tokenu musi być od 4 do 20 znaków.")
    if uczestnicy is not None and len(uczestnicy) != liczba:
        raise ValueError("Liczba uczestników musi odpowiadać liczbie tokenów.")
    uczestnicy = uczestnicy or [(None, None)] * liczba

    # Limit prób, żeby żądanie liczby tokenów przekraczającej pulę możliwych
    # kodów (np. dlugosc=1 i liczba=40) kończyło się czytelnym błędem zamiast
    # zawieszenia aplikacji w nieskończonej pętli (B11).
    maks_prob = max(1000, liczba * 50)

    with baza() as conn:
        istniejace = {w["token"] for w in conn.execute("SELECT token FROM tokeny").fetchall()}
        nowe = []
        proba = 0
        while len(nowe) < liczba:
            if proba >= maks_prob:
                raise ValueError(
                    f"Nie udało się wygenerować {liczba} unikalnych tokenów o długości {dlugosc} — "
                    "zbyt mała pula możliwych kodów dla tej długości. Zwiększ długość tokenu."
                )
            proba += 1
            token = losowy_kod(dlugosc)
            if token not in istniejace and token not in nowe:
                nowe.append(token)

        teraz = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        conn.executemany(
            "INSERT INTO tokeny (token, test_id, wykorzystany, data_utworzenia, imie, email) VALUES (?, ?, 0, ?, ?, ?)",
            [(token, test_id, teraz, *uczestnicy[i]) for i, token in enumerate(nowe)],
        )
    return [{"token": token, "imie": uczestnicy[i][0], "email": uczestnicy[i][1]} for i, token in enumerate(nowe)]


def przypisz_token(token, imie, email):
    """Ustawia/aktualizuje osobę przypisaną do tokenu (już zwalidowane dane).
    Zmiana adresu kasuje stan wysyłki — zaproszenie na nowy adres jeszcze nie
    poszło, więc token wraca do „niewysłanych”. Zwraca True, jeśli token istniał."""
    with baza() as conn:
        cur = conn.execute(
            """
            UPDATE tokeny SET
                mail_wyslany_at = CASE WHEN email IS ? THEN mail_wyslany_at END,
                przypomnienie_wyslane_at = CASE WHEN email IS ? THEN przypomnienie_wyslane_at END,
                mail_status = CASE WHEN email IS ? THEN mail_status END,
                mail_blad = CASE WHEN email IS ? THEN mail_blad END,
                imie = ?, email = ?
            WHERE token = ?
            """,
            (email, email, email, email, imie, email, token),
        )
    return cur.rowcount > 0


# --- Wysyłka maili (Etap 6, F3/E6) ------------------------------------------------------

DOMYSLNE_USTAWIENIA_APLIKACJI = {
    "nazwa_nadawcy": "Test wiedzy",
    "adres_aplikacji": "http://localhost:5555",
}


def pobierz_ustawienia_aplikacji():
    with baza() as conn:
        zapisane = {w["klucz"]: w["wartosc"] for w in conn.execute("SELECT klucz, wartosc FROM ustawienia_aplikacji")}
    return {klucz: zapisane.get(klucz) or domyslna for klucz, domyslna in DOMYSLNE_USTAWIENIA_APLIKACJI.items()}


def zapisz_ustawienia_aplikacji(**ustawienia):
    for klucz in ustawienia:
        if klucz not in DOMYSLNE_USTAWIENIA_APLIKACJI:
            raise ValueError(f"Nieznane ustawienie: {klucz}")
    with baza() as conn:
        conn.executemany(
            "INSERT INTO ustawienia_aplikacji (klucz, wartosc) VALUES (?, ?) "
            "ON CONFLICT(klucz) DO UPDATE SET wartosc = excluded.wartosc",
            list(ustawienia.items()),
        )


def zapisz_szablony_maili(test_id, mail_temat, mail_tresc, przypomnienie_temat, przypomnienie_tresc):
    """Szablony zaproszenia i przypomnienia dla testu — None = szablon domyślny."""
    with baza() as conn:
        cur = conn.execute(
            "UPDATE testy SET mail_temat = ?, mail_tresc = ?, przypomnienie_temat = ?, przypomnienie_tresc = ? "
            "WHERE id = ?",
            (mail_temat, mail_tresc, przypomnienie_temat, przypomnienie_tresc, test_id),
        )
    return cur.rowcount > 0


RODZAJE_WYSYLKI = {"zaproszenia", "przypomnienia", "zaznaczone"}


def tokeny_do_wysylki(test_id, rodzaj, zaznaczone=None):
    """Tokeny z adresem e-mail, do których ma pójść mail:
    - `zaproszenia` — zaproszenie jeszcze nie zostało wysłane (także po błędzie),
    - `przypomnienia` (E6) — osoba nie ukończyła testu (nie zaczęła albo w trakcie),
    - `zaznaczone` — zaproszenie (ponownie) do tokenów wybranych na liście."""
    if rodzaj not in RODZAJE_WYSYLKI:
        raise ValueError(f"Nieznany rodzaj wysyłki: {rodzaj}")
    zapytanie = """
        SELECT tk.token, tk.imie, tk.email
        FROM tokeny tk
        LEFT JOIN podejscia pj ON pj.token = tk.token
        WHERE tk.test_id = ? AND tk.email IS NOT NULL
    """
    parametry = [test_id]
    if rodzaj == "zaproszenia":
        # Bez tokenów już ukończonych — np. rozdanych wcześniej ręcznie.
        zapytanie += " AND tk.mail_wyslany_at IS NULL AND COALESCE(pj.status, 'wolny') IN ('wolny', 'w_trakcie')"
    elif rodzaj == "przypomnienia":
        zapytanie += " AND COALESCE(pj.status, 'wolny') IN ('wolny', 'w_trakcie')"
    else:
        zaznaczone = list(zaznaczone or [])
        if not zaznaczone:
            return []
        zapytanie += f" AND tk.token IN ({', '.join('?' for _ in zaznaczone)})"
        parametry += zaznaczone
    zapytanie += " ORDER BY tk.data_utworzenia, tk.token"
    with baza() as conn:
        return [dict(w) for w in conn.execute(zapytanie, parametry).fetchall()]


def zapisz_wynik_wysylki(token, przypomnienie, blad=None):
    """Wynik wysłania jednego maila: `mail_status`/`mail_blad` opisują ostatnią
    próbę (dowolnego rodzaju), data trafia do kolumny zaproszenia albo
    przypomnienia tylko przy sukcesie."""
    teraz = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    kolumna_daty = "przypomnienie_wyslane_at" if przypomnienie else "mail_wyslany_at"
    with baza() as conn:
        if blad is None:
            conn.execute(
                f"UPDATE tokeny SET mail_status = 'ok', mail_blad = NULL, {kolumna_daty} = ? WHERE token = ?",
                (teraz, token),
            )
        else:
            conn.execute("UPDATE tokeny SET mail_status = 'blad', mail_blad = ? WHERE token = ?", (blad, token))


def resetuj_token(token):
    """Akcja admina (L2) — kasuje podejście i dotychczasowe odpowiedzi powiązane
    z tokenem i odblokowuje go, żeby uczestnik mógł zacząć test od nowa.
    Zgodnie z decyzją z Etapu 2: reset usuwa poprzednie odpowiedzi (nie
    archiwizuje ich jako osobne podejście). Zwraca True, jeśli token istniał."""
    with baza() as conn:
        token_wiersz = conn.execute("SELECT 1 FROM tokeny WHERE token = ?", (token,)).fetchone()
        if token_wiersz is None:
            return False
        conn.execute("DELETE FROM odpowiedzi_uzytkownika WHERE token = ?", (token,))
        conn.execute("DELETE FROM podejscia WHERE token = ?", (token,))
        conn.execute(
            "UPDATE tokeny SET wykorzystany = 0, data_wykorzystania = NULL WHERE token = ?",
            (token,),
        )
    return True


TRYBY_SZCZEGOLOW = {"natychmiast", "po_zamknieciu", "od_daty", "nigdy"}
TRYBY_WYNIKU = {"od_razu", "razem_ze_szczegolami"}
FORMAT_DATY = "%Y-%m-%d %H:%M:%S"


# Znacznik „nie zmieniaj” dla parametrów, w których None ma własne znaczenie
# (np. prog_zaliczenia=None to „bez progu”, a nie „zostaw jak było”).
BEZ_ZMIAN = object()


def zapisz_ustawienia_testu(
    test_id, limit_czasu_min, tryb_szczegolow, szczegoly_od, wynik_widoczny, dostepny_od, dostepny_do,
    liczba_pytan_do_losowania=None, prog_zaliczenia=BEZ_ZMIAN,
):
    """Aktualizuje ustawienia kontroli nad testem (Etap 4, F1/F2). Zmiana
    działa też na trwające podejścia, bo ustawienia są czytane na żywo przy
    każdym wejściu na `/test` — zgodnie z decyzją usera z Etapu 4.
    `liczba_pytan_do_losowania=None` zostawia dotychczasową wartość bez zmian
    (np. wywołania z testów, które jej nie dotyczą). `prog_zaliczenia` (E2,
    Etap 5) to procent 1–100 albo None (bez progu); pominięty — bez zmian."""
    if tryb_szczegolow not in TRYBY_SZCZEGOLOW:
        raise ValueError(f"Nieprawidłowy tryb szczegółów: {tryb_szczegolow}")
    if wynik_widoczny not in TRYBY_WYNIKU:
        raise ValueError(f"Nieprawidłowy tryb widoczności wyniku: {wynik_widoczny}")
    if prog_zaliczenia is not BEZ_ZMIAN and prog_zaliczenia is not None and not 1 <= prog_zaliczenia <= 100:
        raise ValueError("Próg zaliczenia musi być od 1 do 100%.")

    kolumny = ["limit_czasu_min", "tryb_szczegolow", "szczegoly_od", "wynik_widoczny", "dostepny_od", "dostepny_do"]
    wartosci = [limit_czasu_min, tryb_szczegolow, szczegoly_od, wynik_widoczny, dostepny_od, dostepny_do]
    if liczba_pytan_do_losowania is not None:
        kolumny.append("liczba_pytan_do_losowania")
        wartosci.append(liczba_pytan_do_losowania)
    if prog_zaliczenia is not BEZ_ZMIAN:
        kolumny.append("prog_zaliczenia")
        wartosci.append(prog_zaliczenia)

    # Nazwy kolumn pochodzą wyłącznie z powyższej stałej listy, nie z wejścia.
    with baza() as conn:
        cur = conn.execute(
            f"UPDATE testy SET {', '.join(f'{k} = ?' for k in kolumny)} WHERE id = ?",
            (*wartosci, test_id),
        )
    return cur.rowcount > 0


def ustaw_zamkniecie_testu(test_id, zamkniety):
    """„Zamknij test i opublikuj odpowiedzi” / ponowne otwarcie (F2, punkt 5) —
    zamknięcie blokuje nowe podejścia (start.html), a w trybie szczegółów
    `po_zamknieciu` odblokowuje uczestnikom poprawne odpowiedzi."""
    with baza() as conn:
        cur = conn.execute("UPDATE testy SET zamkniety = ? WHERE id = ?", (1 if zamkniety else 0, test_id))
    return cur.rowcount > 0


def oblicz_termin(data_rozpoczecia_str, limit_czasu_min):
    """Nominalny termin upłynięcia limitu czasu (bez tolerancji) — używane do
    wyświetlania licznika uczestnikowi (F1)."""
    start = datetime.strptime(data_rozpoczecia_str, FORMAT_DATY)
    return start + timedelta(minutes=limit_czasu_min)


def domknij_przeterminowane_podejscia(test_id=None):
    """Leniwe domykanie podejść, których limit czasu (+ 5 s tolerancji) minął,
    a status wciąż jest `w_trakcie` (F1, punkt 5) — wywoływane przy wejściu
    uczestnika na `/test`/`/wynik` i przy wyświetlaniu list w panelu admina,
    żeby statusy tam były aktualne bez osobnego zadania w tle."""
    teraz = datetime.now()
    with baza() as conn:
        zapytanie = """
            SELECT pj.token AS token, pj.data_rozpoczecia AS data_rozpoczecia, t.limit_czasu_min AS limit_czasu_min
            FROM podejscia pj
            JOIN testy t ON t.id = pj.test_id
            WHERE pj.status = 'w_trakcie' AND t.limit_czasu_min IS NOT NULL
        """
        parametry = ()
        if test_id is not None:
            zapytanie += " AND pj.test_id = ?"
            parametry = (test_id,)
        kandydaci = conn.execute(zapytanie, parametry).fetchall()

        for wiersz in kandydaci:
            termin = oblicz_termin(wiersz["data_rozpoczecia"], wiersz["limit_czasu_min"]) + timedelta(seconds=TOLERANCJA_SEKUNDY)
            if teraz > termin:
                conn.execute(
                    "UPDATE podejscia SET status = 'czas_minal', data_zakonczenia = ? WHERE token = ? AND status = 'w_trakcie'",
                    (teraz.strftime(FORMAT_DATY), wiersz["token"]),
                )


def czy_zaliczony(procent, prog_zaliczenia):
    """Zdał/nie zdał (E2) — None, gdy test nie ma ustawionego progu."""
    if prog_zaliczenia is None:
        return None
    return procent >= prog_zaliczenia


def szczegoly_podejsc_testu(test_id, token=None):
    """Dane do szczegółowego eksportu (E1) — dla każdego zakończonego podejścia
    wszystkie wylosowane pytania w kolejności wyświetlania, także te bez
    odpowiedzi (np. gdy minął czas), których nie ma w widoku arkusz_wynikow.
    Z `token` — tylko podejście jednej osoby (eksport indywidualny, Etap 6)."""
    zapytanie = """
        SELECT pj.token AS token, tk.imie AS imie, tk.email AS email, pj.wybrane_pytania AS wybrane_pytania
        FROM podejscia pj
        JOIN tokeny tk ON tk.token = pj.token
        WHERE pj.test_id = ? AND pj.status IN ('zakonczone', 'czas_minal')
    """
    parametry = [test_id]
    if token is not None:
        zapytanie += " AND pj.token = ?"
        parametry.append(token)
    with baza() as conn:
        podejscia = conn.execute(zapytanie + " ORDER BY pj.data_zakonczenia, pj.token", parametry).fetchall()
        pytania = {
            w["id"]: w
            for w in conn.execute("SELECT * FROM pytania WHERE test_id = ?", (test_id,)).fetchall()
        }
        odpowiedzi = {
            (w["token"], w["pytanie_id"]): w["odpowiedz"]
            for w in conn.execute(
                """
                SELECT ou.token, ou.pytanie_id, ou.odpowiedz
                FROM odpowiedzi_uzytkownika ou JOIN podejscia pj ON pj.token = ou.token
                WHERE pj.test_id = ?
                """,
                (test_id,),
            ).fetchall()
        }

    wynik = []
    for podejscie in podejscia:
        for numer, pytanie_id in enumerate(json.loads(podejscie["wybrane_pytania"]), start=1):
            pytanie = pytania.get(pytanie_id)
            if pytanie is None:
                continue
            dana = odpowiedzi.get((podejscie["token"], pytanie_id))
            wynik.append(
                {
                    "token": podejscie["token"],
                    "imie": podejscie["imie"],
                    "email": podejscie["email"],
                    "numer": numer,
                    "tresc_pytania": pytanie["tresc_pytania"],
                    "odpowiedz_uzytkownika": dana,
                    "tresc_odpowiedzi_uzytkownika": pytanie[f"opcja_{dana.lower()}"] if dana else None,
                    "poprawna_odpowiedz": pytanie["odpowiedz"],
                    "tresc_poprawnej_odpowiedzi": pytanie[f"opcja_{pytanie['odpowiedz'].lower()}"],
                    "czy_poprawna": dana == pytanie["odpowiedz"],
                    "wyjasnienie": pytanie["wyjasnienie"],
                }
            )
    return wynik


def statystyki_pytan(test_id):
    """Statystyki per pytanie z banku (E3) na podstawie zakończonych podejść:
    ile razy pytanie wylosowano, ile było poprawnych odpowiedzi, rozkład
    wyborów A–D i brak odpowiedzi. Procent poprawnych liczony od liczby
    wylosowań — pytanie bez odpowiedzi liczy się jako błędne, tak jak w
    wyniku uczestnika (L3). Posortowane od najsłabiej rozwiązywanych, bo po
    to są te statystyki: pytanie z bardzo niskim wynikiem to często zła litera
    w pliku albo temat do powtórzenia na szkoleniu."""
    with baza() as conn:
        pytania = conn.execute("SELECT * FROM pytania WHERE test_id = ? ORDER BY id", (test_id,)).fetchall()
        podejscia = conn.execute(
            "SELECT wybrane_pytania FROM podejscia WHERE test_id = ? AND status IN ('zakonczone', 'czas_minal')",
            (test_id,),
        ).fetchall()
        odpowiedzi = conn.execute(
            """
            SELECT ou.pytanie_id AS pytanie_id, ou.odpowiedz AS odpowiedz, COUNT(*) AS ile
            FROM odpowiedzi_uzytkownika ou JOIN podejscia pj ON pj.token = ou.token
            WHERE pj.test_id = ? AND pj.status IN ('zakonczone', 'czas_minal')
            GROUP BY ou.pytanie_id, ou.odpowiedz
            """,
            (test_id,),
        ).fetchall()

    wylosowania = {}
    for podejscie in podejscia:
        for pytanie_id in json.loads(podejscie["wybrane_pytania"]):
            wylosowania[pytanie_id] = wylosowania.get(pytanie_id, 0) + 1

    rozklady = {}
    for w in odpowiedzi:
        rozklady.setdefault(w["pytanie_id"], {})[w["odpowiedz"]] = w["ile"]

    statystyki = []
    for pytanie in pytania:
        ile_razy = wylosowania.get(pytanie["id"], 0)
        rozklad = {litera: rozklady.get(pytanie["id"], {}).get(litera, 0) for litera in "ABCD"}
        poprawne = rozklad[pytanie["odpowiedz"]]
        statystyki.append(
            {
                "tresc_pytania": pytanie["tresc_pytania"],
                "poprawna_odpowiedz": pytanie["odpowiedz"],
                "wylosowane": ile_razy,
                "poprawne": poprawne,
                "procent": round(100 * poprawne / ile_razy, 1) if ile_razy else None,
                "rozklad": rozklad,
                "bez_odpowiedzi": max(0, ile_razy - sum(rozklad.values())),
            }
        )
    # Pytania jeszcze nigdy niewylosowane (procent None) na końcu listy.
    statystyki.sort(key=lambda s: (s["procent"] is None, s["procent"] or 0))
    return statystyki


def kopia_bazy():
    """Spójna kopia całej bazy jako bajty (S7) — `backup()` z SQLite działa
    poprawnie także w trakcie zapisów (WAL), w przeciwieństwie do zwykłego
    skopiowania pliku baza.db, który może nie zawierać zmian z baza.db-wal."""
    zrodlo = polacz()
    kopia = sqlite3.connect(":memory:")
    try:
        zrodlo.backup(kopia)
        return kopia.serialize()
    finally:
        kopia.close()
        zrodlo.close()
