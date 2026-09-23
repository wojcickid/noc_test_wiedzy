"""Testy Etapu 5 (patrz PLAN_POPRAWEK.md): eksport wyników do XLSX (E1/UI10),
kopiowanie/eksport tokenów (UI9), próg zaliczenia (E2), statystyki pytań (E3)
i kopia bazy z panelu (S7)."""

import csv
import io
import sqlite3
from datetime import datetime, timedelta

import openpyxl

import db
from conftest import pobierz_csrf_token, zaloguj_admina
from pomocnicze import przejdz_caly_test


def _ukonczony_token(klient, test_id, litera="B", przypisany=None):
    """Generuje token i przechodzi nim cały test, odpowiadając `litera`."""
    (wpis,) = db.wygeneruj_tokeny(test_id, 1, dlugosc=6, uczestnicy=[(przypisany, None)] if przypisany else None)
    przejdz_caly_test(klient, wpis["token"], litera)
    return wpis["token"]


def _ustaw_prog(test_id, prog):
    with db.baza() as conn:
        conn.execute("UPDATE testy SET prog_zaliczenia = ? WHERE id = ?", (prog, test_id))


def _xlsx(odpowiedz):
    return openpyxl.load_workbook(io.BytesIO(odpowiedz.data))


# --- E2: próg zaliczenia ----------------------------------------------------------------

def test_nowy_test_ma_domyslny_prog_80(klient, test_z_pytaniami):
    with db.baza() as conn:
        prog = conn.execute("SELECT prog_zaliczenia FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()[0]
    assert prog == 80


def test_migracja_dodaje_prog_80_do_istniejacych_testow(tmp_path, monkeypatch):
    """Baza sprzed Etapu 5 (bez kolumny prog_zaliczenia) po starcie dostaje
    kolumnę z wartością 80 dla już istniejących testów, bez utraty danych."""
    plik = tmp_path / "stara.db"
    monkeypatch.setattr(db, "DB_PLIK", str(plik))
    db.inicjalizuj()
    db.importuj_pytania("Stary test", [
        {"tresc_pytania": "P?", "opcja_a": "a", "opcja_b": "b", "opcja_c": "c", "opcja_d": "d", "odpowiedz": "A"},
    ], liczba_pytan=1)
    conn = sqlite3.connect(plik)
    conn.executescript("DROP VIEW zbiorcze_wyniki; DROP VIEW arkusz_wynikow; ALTER TABLE testy DROP COLUMN prog_zaliczenia;")
    conn.close()

    db.inicjalizuj()
    db.inicjalizuj()  # idempotentnie

    with db.baza() as conn:
        wiersz = conn.execute("SELECT nazwa, prog_zaliczenia FROM testy").fetchone()
    assert (wiersz["nazwa"], wiersz["prog_zaliczenia"]) == ("Stary test", 80)


def test_wynik_pokazuje_zaliczenie(klient, test_z_pytaniami):
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    strona = przejdz_caly_test(klient, token, "B")
    tekst = strona.data.decode()
    assert "Test zaliczony" in tekst
    assert "próg zaliczenia: 80%" in tekst


def test_wynik_pokazuje_niezaliczenie(klient, test_z_pytaniami):
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    strona = przejdz_caly_test(klient, token, "A")
    assert "Test niezaliczony" in strona.data.decode()


def test_bez_progu_brak_informacji_o_zaliczeniu(klient, test_z_pytaniami):
    _ustaw_prog(test_z_pytaniami, None)
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    strona = przejdz_caly_test(klient, token, "B")
    tekst = strona.data.decode()
    assert "zaliczony" not in tekst
    assert "próg zaliczenia" not in tekst


def test_prog_rowny_wynikowi_zalicza():
    assert db.czy_zaliczony(80.0, 80) is True
    assert db.czy_zaliczony(79.9, 80) is False
    assert db.czy_zaliczony(10, None) is None


def test_zaliczenie_ukryte_razem_z_wynikiem(klient, test_z_pytaniami):
    """Zdał/nie zdał nie może wyciec wcześniej niż sam wynik punktowy (decyzja z Etapu 5)."""
    db.zapisz_ustawienia_testu(test_z_pytaniami, None, "nigdy", None, "razem_ze_szczegolami", None, None)
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    klient.post("/", data={"token": token})
    klient.post("/test", data={"start": "1"})
    with db.baza() as conn:
        ids = [w["id"] for w in conn.execute("SELECT id FROM pytania WHERE test_id = ?", (test_z_pytaniami,))]
    for pytanie_id in ids:
        strona = klient.post("/test", data={"pytanie_id": str(pytanie_id), "odpowiedz": "B"}, follow_redirects=True)
    assert "zaliczony" not in strona.data.decode()

    strona = klient.post("/sprawdz-wynik", data={"token": token})
    assert "zaliczony" not in strona.data.decode()


def test_sprawdz_wynik_liczy_zaliczenie_od_liczby_wylosowanych_pytan(klient, test_z_pytaniami):
    """Przy „czas minął” uczestnik odpowiedział na 1 z 3 pytań (poprawnie) —
    wynik to 1/3 (33%), nie 1/1 (100%), więc test jest niezaliczony."""
    db.zapisz_ustawienia_testu(test_z_pytaniami, 1, "natychmiast", None, "od_razu", None, None)
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    klient.post("/", data={"token": token})
    strona = klient.post("/test", data={"start": "1"}, follow_redirects=True)
    import re
    pytanie_id = re.search(rb'name="pytanie_id" value="(\d+)"', strona.data).group(1).decode()
    klient.post("/test", data={"pytanie_id": pytanie_id, "odpowiedz": "B"})
    stara = (datetime.now() - timedelta(minutes=5)).strftime(db.FORMAT_DATY)
    with db.baza() as conn:
        conn.execute("UPDATE podejscia SET data_rozpoczecia = ? WHERE token = ?", (stara, token))

    strona = klient.post("/sprawdz-wynik", data={"token": token})
    tekst = strona.data.decode()
    assert "1 / 3" in tekst
    assert "Test niezaliczony" in tekst


def test_ustawienia_zapisuja_i_zeruja_prog(klient, test_z_pytaniami, haslo_admina):
    zaloguj_admina(klient, haslo_admina)
    klient.get(f"/admin/testy/{test_z_pytaniami}/ustawienia")
    dane = {
        "csrf_token": pobierz_csrf_token(klient),
        "tryb_szczegolow": "po_zamknieciu",
        "wynik_widoczny": "od_razu",
        "prog_zaliczenia": "65",
    }
    klient.post(f"/admin/testy/{test_z_pytaniami}/ustawienia", data=dane)
    with db.baza() as conn:
        assert conn.execute("SELECT prog_zaliczenia FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()[0] == 65

    dane["prog_zaliczenia"] = ""
    klient.post(f"/admin/testy/{test_z_pytaniami}/ustawienia", data=dane)
    with db.baza() as conn:
        assert conn.execute("SELECT prog_zaliczenia FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()[0] is None


def test_ustawienia_odrzucaja_nieprawidlowy_prog(klient, test_z_pytaniami, haslo_admina):
    zaloguj_admina(klient, haslo_admina)
    klient.get(f"/admin/testy/{test_z_pytaniami}/ustawienia")
    for zly in ["0", "101", "abc"]:
        strona = klient.post(
            f"/admin/testy/{test_z_pytaniami}/ustawienia",
            data={
                "csrf_token": pobierz_csrf_token(klient),
                "tryb_szczegolow": "po_zamknieciu",
                "wynik_widoczny": "od_razu",
                "prog_zaliczenia": zly,
            },
            follow_redirects=True,
        )
        assert "Próg zaliczenia musi być" in strona.data.decode()
    with db.baza() as conn:
        assert conn.execute("SELECT prog_zaliczenia FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()[0] == 80


def test_panel_testu_pokazuje_kolumne_zaliczenia(klient, test_z_pytaniami, haslo_admina):
    _ukonczony_token(klient, test_z_pytaniami, "B")
    _ukonczony_token(klient, test_z_pytaniami, "A")
    zaloguj_admina(klient, haslo_admina)
    tekst = klient.get(f"/admin/testy/{test_z_pytaniami}").data.decode()
    assert "<th>Zaliczenie</th>" in tekst
    assert ">zdał<" in tekst
    assert ">nie zdał<" in tekst


# --- E1/UI10: eksport wyników do XLSX -------------------------------------------------

def test_eksport_wynikow_wymaga_logowania(klient, test_z_pytaniami):
    odpowiedz = klient.get(f"/admin/testy/{test_z_pytaniami}/wyniki.xlsx")
    assert odpowiedz.status_code == 302
    assert "/admin/login" in odpowiedz.headers["Location"]


def test_eksport_wynikow_arkusz_zbiorczy(klient, test_z_pytaniami, haslo_admina):
    token = _ukonczony_token(klient, test_z_pytaniami, "B", przypisany="Zażółć Gęślą")
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get(f"/admin/testy/{test_z_pytaniami}/wyniki.xlsx")
    assert odpowiedz.status_code == 200
    assert "Test_testowy_wyniki.xlsx" in odpowiedz.headers["Content-Disposition"]

    skoroszyt = _xlsx(odpowiedz)
    assert skoroszyt.sheetnames == ["Zbiorczo", "Szczegółowo"]
    wiersze = list(skoroszyt["Zbiorczo"].iter_rows(values_only=True))
    assert wiersze[0][-1] == "Zaliczenie (próg 80%)"
    wiersz = wiersze[1]
    assert wiersz[0] == token
    assert wiersz[1] == "Zażółć Gęślą"
    assert wiersz[2] is None
    assert wiersz[3] == "ukończony"
    assert isinstance(wiersz[4], datetime) and isinstance(wiersz[5], datetime)
    assert wiersz[7:10] == (3, 3, 1.0)
    assert wiersz[10] == "zdał"


def test_eksport_wynikow_bez_progu_nie_ma_kolumny_zaliczenia(klient, test_z_pytaniami, haslo_admina):
    _ustaw_prog(test_z_pytaniami, None)
    _ukonczony_token(klient, test_z_pytaniami, "B")
    zaloguj_admina(klient, haslo_admina)
    arkusz = _xlsx(klient.get(f"/admin/testy/{test_z_pytaniami}/wyniki.xlsx"))["Zbiorczo"]
    assert arkusz.cell(1, arkusz.max_column).value == "Procent"


def test_eksport_wynikow_arkusz_szczegolowy_z_brakiem_odpowiedzi(klient, test_z_pytaniami, haslo_admina):
    """Pytania bez odpowiedzi (czas minął) też są w arkuszu szczegółowym."""
    db.zapisz_ustawienia_testu(test_z_pytaniami, 1, "po_zamknieciu", None, "od_razu", None, None)
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    klient.post("/", data={"token": token})
    klient.post("/test", data={"start": "1"})
    stara = (datetime.now() - timedelta(minutes=5)).strftime(db.FORMAT_DATY)
    with db.baza() as conn:
        conn.execute("UPDATE podejscia SET data_rozpoczecia = ? WHERE token = ?", (stara, token))

    zaloguj_admina(klient, haslo_admina)
    skoroszyt = _xlsx(klient.get(f"/admin/testy/{test_z_pytaniami}/wyniki.xlsx"))
    zbiorczo = list(skoroszyt["Zbiorczo"].iter_rows(values_only=True))
    assert zbiorczo[1][3] == "czas minął"
    szczegoly = list(skoroszyt["Szczegółowo"].iter_rows(values_only=True))[1:]
    assert [w[3] for w in szczegoly] == [1, 2, 3]
    assert all(w[5] == "brak odpowiedzi" and w[7] == "błędna" for w in szczegoly)


def test_eksport_wynikow_nie_zawiera_podejsc_w_trakcie(klient, test_z_pytaniami, haslo_admina):
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    klient.post("/", data={"token": token})
    klient.post("/test", data={"start": "1"})
    zaloguj_admina(klient, haslo_admina)
    skoroszyt = _xlsx(klient.get(f"/admin/testy/{test_z_pytaniami}/wyniki.xlsx"))
    assert skoroszyt["Zbiorczo"].max_row == 1
    assert skoroszyt["Szczegółowo"].max_row == 1


def test_eksport_wynikow_nieistniejacego_testu(klient, haslo_admina):
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get("/admin/testy/999/wyniki.xlsx")
    assert odpowiedz.status_code == 302


# --- UI9: tokeny — kopiowanie i CSV ---------------------------------------------------

def test_eksport_tokenow_csv(klient, test_z_pytaniami, haslo_admina):
    db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Łukasz Żółw", "lukasz@firma.pl")])
    ukonczony = _ukonczony_token(klient, test_z_pytaniami, "B")
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get(f"/admin/testy/{test_z_pytaniami}/tokeny.csv")
    assert odpowiedz.status_code == 200
    assert odpowiedz.data.startswith("﻿".encode("utf-8"))

    wiersze = list(csv.reader(io.StringIO(odpowiedz.data.decode("utf-8-sig")), delimiter=";"))
    assert wiersze[0] == ["token", "imie", "email", "status"]
    statusy = {w[0]: (w[1], w[2], w[3]) for w in wiersze[1:]}
    assert statusy[ukonczony] == ("", "", "ukończony")
    assert ("Łukasz Żółw", "lukasz@firma.pl", "wolny") in statusy.values()


def test_eksport_tokenow_wymaga_logowania(klient, test_z_pytaniami):
    assert klient.get(f"/admin/testy/{test_z_pytaniami}/tokeny.csv").status_code == 302


def test_strona_tokenow_ma_tekst_do_skopiowania(klient, test_z_pytaniami, haslo_admina):
    (wpis,) = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Anna Nowak", None)])
    zaloguj_admina(klient, haslo_admina)
    tekst = klient.get(f"/admin/testy/{test_z_pytaniami}/tokeny").data.decode()
    assert "Kopiuj wszystkie" in tekst
    assert f"Anna Nowak\t\t{wpis['token']}" in tekst


def test_panel_ma_link_do_tokenow_w_wierszu_testu(klient, test_z_pytaniami, haslo_admina):
    zaloguj_admina(klient, haslo_admina)
    tekst = klient.get("/admin").data.decode()
    assert f'href="/admin/testy/{test_z_pytaniami}/tokeny"' in tekst


# --- E3: statystyki pytań -------------------------------------------------------------

def test_statystyki_pytan(klient, test_z_pytaniami, haslo_admina):
    _ukonczony_token(klient, test_z_pytaniami, "B")
    _ukonczony_token(klient, test_z_pytaniami, "A")
    _ukonczony_token(klient, test_z_pytaniami, "B")

    statystyki = db.statystyki_pytan(test_z_pytaniami)
    assert len(statystyki) == 3
    for s in statystyki:
        assert s["wylosowane"] == 3
        assert s["poprawne"] == 2
        assert s["procent"] == 66.7
        assert s["rozklad"] == {"A": 1, "B": 2, "C": 0, "D": 0}
        assert s["bez_odpowiedzi"] == 0

    zaloguj_admina(klient, haslo_admina)
    strona = klient.get(f"/admin/testy/{test_z_pytaniami}/statystyki")
    assert strona.status_code == 200
    assert "66.7%" in strona.data.decode()


def test_statystyki_sortuja_od_najslabszych_i_licza_brak_odpowiedzi(klient, haslo_admina):
    pytania = [
        {"tresc_pytania": "Łatwe", "opcja_a": "a", "opcja_b": "b", "opcja_c": "c", "opcja_d": "d", "odpowiedz": "A"},
        {"tresc_pytania": "Trudne", "opcja_a": "a", "opcja_b": "b", "opcja_c": "c", "opcja_d": "d", "odpowiedz": "D"},
        {"tresc_pytania": "Nieużyte", "opcja_a": "a", "opcja_b": "b", "opcja_c": "c", "opcja_d": "d", "odpowiedz": "A"},
    ]
    test_id, _, _ = db.importuj_pytania("Statystyki", pytania, liczba_pytan=3)
    with db.baza() as conn:
        ids = {w["tresc_pytania"]: w["id"] for w in conn.execute("SELECT * FROM pytania WHERE test_id = ?", (test_id,))}
        # Podejście ręczne: wylosowane tylko „Łatwe” i „Trudne”, na „Trudne” brak odpowiedzi.
        db.wygeneruj_tokeny(test_id, 1, dlugosc=6, uczestnicy=[("x", None)])
        token = conn.execute("SELECT token FROM tokeny WHERE test_id = ?", (test_id,)).fetchone()[0]
        conn.execute(
            "INSERT INTO podejscia (token, test_id, wybrane_pytania, liczba_pytan, indeks_pytania, status, data_rozpoczecia, data_zakonczenia) "
            "VALUES (?, ?, ?, 2, 1, 'czas_minal', '2026-01-01 10:00:00', '2026-01-01 10:05:00')",
            (token, test_id, f"[{ids['Łatwe']}, {ids['Trudne']}]"),
        )
        conn.execute(
            "INSERT INTO odpowiedzi_uzytkownika (token, pytanie_id, odpowiedz, data_wyslania) VALUES (?, ?, 'A', '2026-01-01 10:01:00')",
            (token, ids["Łatwe"]),
        )

    statystyki = db.statystyki_pytan(test_id)
    assert [s["tresc_pytania"] for s in statystyki] == ["Trudne", "Łatwe", "Nieużyte"]
    trudne = statystyki[0]
    assert (trudne["wylosowane"], trudne["procent"], trudne["bez_odpowiedzi"]) == (1, 0.0, 1)
    assert statystyki[2]["procent"] is None


def test_statystyki_wymagaja_logowania(klient, test_z_pytaniami):
    assert klient.get(f"/admin/testy/{test_z_pytaniami}/statystyki").status_code == 302


# --- S7: kopia bazy -------------------------------------------------------------------

def test_kopia_bazy_do_pobrania(klient, test_z_pytaniami, haslo_admina, tmp_path):
    token = _ukonczony_token(klient, test_z_pytaniami, "B")
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get("/admin/kopia-bazy")
    assert odpowiedz.status_code == 200
    assert "baza-" in odpowiedz.headers["Content-Disposition"]
    assert odpowiedz.data.startswith(b"SQLite format 3")

    plik = tmp_path / "kopia.db"
    plik.write_bytes(odpowiedz.data)
    conn = sqlite3.connect(plik)
    try:
        assert conn.execute("SELECT COUNT(*) FROM pytania").fetchone()[0] == 3
        assert conn.execute("SELECT status FROM podejscia WHERE token = ?", (token,)).fetchone()[0] == "zakonczone"
    finally:
        conn.close()


def test_kopia_bazy_wymaga_logowania(klient):
    odpowiedz = klient.get("/admin/kopia-bazy")
    assert odpowiedz.status_code == 302
    assert "/admin/login" in odpowiedz.headers["Location"]
