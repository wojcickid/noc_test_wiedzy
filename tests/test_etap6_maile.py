"""Testy Etapu 6 (patrz PLAN_POPRAWEK.md): imię i e-mail zamiast jednego pola
`przypisany` (L6), wysyłka maili przez SMTP (F3), przypomnienia (E6), eksport
do korespondencji seryjnej, wynik jednej osoby do XLSX i token z linku.
Serwer SMTP jest podmieniony na atrapę — testy niczego nie wysyłają."""

import io
import json
import re
import smtplib
import sqlite3
from datetime import datetime, timedelta

import openpyxl
import pytest

import db
import poczta
import test_wiedzy_app as app_module
from conftest import pobierz_csrf_token, zaloguj_admina
from pomocnicze import przejdz_caly_test, zbuduj_xlsx

HASLO_SMTP = "SEKRETNE-haslo-aplikacji-123"


class AtrapaSMTP:
    """Zamiast smtplib.SMTP — zapamiętuje wiadomości, umie udawać błędy."""

    wyslane = []
    odrzucane_adresy = set()
    zle_haslo = False

    def __init__(self, host, port, timeout=None):
        self.host, self.port = host, port

    def starttls(self, context=None):
        pass

    def login(self, login, haslo):
        if AtrapaSMTP.zle_haslo:
            raise smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted")

    def send_message(self, wiadomosc):
        if wiadomosc["To"] in AtrapaSMTP.odrzucane_adresy:
            raise smtplib.SMTPRecipientsRefused({wiadomosc["To"]: (550, b"5.1.1 No such user")})
        AtrapaSMTP.wyslane.append(wiadomosc)

    def quit(self):
        pass


@pytest.fixture
def smtp(tmp_path, monkeypatch):
    """config.json w katalogu tymczasowym + atrapa serwera, wysyłka synchroniczna."""
    plik = tmp_path / "config.json"
    plik.write_text(json.dumps({"smtp": {
        "host": "smtp.example.com", "port": 587, "szyfrowanie": "starttls",
        "login": "nadawca@example.com", "haslo": HASLO_SMTP,
    }}), encoding="utf-8")
    monkeypatch.setattr(poczta, "CONFIG_PLIK", str(plik))
    monkeypatch.setattr(poczta, "ODSTEP_SEKUNDY", 0)
    monkeypatch.setattr(poczta, "_wysylki", {})
    monkeypatch.setattr(poczta.smtplib, "SMTP", AtrapaSMTP)
    monkeypatch.setattr(AtrapaSMTP, "wyslane", [])
    monkeypatch.setattr(AtrapaSMTP, "odrzucane_adresy", set())
    monkeypatch.setattr(AtrapaSMTP, "zle_haslo", False)
    monkeypatch.setitem(app_module.app.config, "WYSYLKA_W_TLE", False)
    return AtrapaSMTP


@pytest.fixture
def bez_konfiguracji(tmp_path, monkeypatch):
    monkeypatch.setattr(poczta, "CONFIG_PLIK", str(tmp_path / "brak.json"))


def _tokeny(test_id):
    with db.baza() as conn:
        return {w["token"]: dict(w) for w in conn.execute("SELECT * FROM tokeny WHERE test_id = ?", (test_id,))}


def _wyslij(klient, test_id, rodzaj, tokeny=()):
    return klient.post(
        f"/admin/testy/{test_id}/wyslij",
        data={"rodzaj": rodzaj, "tokeny": list(tokeny), "csrf_token": pobierz_csrf_token(klient)},
        follow_redirects=True,
    )


def _xlsx(odpowiedz):
    return openpyxl.load_workbook(io.BytesIO(odpowiedz.data))


# --- L6: lista uczestników ------------------------------------------------------------

def test_parsowanie_listy_uczestnikow_rozne_formaty():
    tekst = (
        "Jan Kowalski;jan.kowalski@firma.pl\n"
        "Anna Nowak <anna.nowak@firma.pl>\n"
        "Piotr Wiśniewski\tpiotr@firma.pl\n"
        "sam.adres@firma.pl\n"
        "\n"
        "Zenon Bez Maila\n"
    )
    uczestnicy, bledy = db.parsuj_liste_uczestnikow(tekst)
    assert bledy == []
    assert uczestnicy == [
        ("Jan Kowalski", "jan.kowalski@firma.pl"),
        ("Anna Nowak", "anna.nowak@firma.pl"),
        ("Piotr Wiśniewski", "piotr@firma.pl"),
        (None, "sam.adres@firma.pl"),
        ("Zenon Bez Maila", None),
    ]


def test_parsowanie_listy_uczestnikow_bledy():
    tekst = (
        "Jan;jan@@firma\n"
        "Anna;anna@firma.pl\n"
        "Druga Anna;ANNA@firma.pl\n"
        "Dwa Adresy a@firma.pl b@firma.pl\n"
    )
    _, bledy = db.parsuj_liste_uczestnikow(tekst)
    tekst_bledow = "\n".join(bledy)
    assert "linia 1: nieprawidłowy adres e-mail" in tekst_bledow
    assert "linia 3: adres" in tekst_bledow and "(pierwszy raz: linia 2)" in tekst_bledow
    assert "linia 4: więcej niż jeden adres e-mail" in tekst_bledow


def test_parsowanie_listy_uczestnikow_limit():
    tekst = "\n".join(f"osoba{i}@firma.pl" for i in range(db.MAKS_UCZESTNIKOW + 1))
    _, bledy = db.parsuj_liste_uczestnikow(tekst)
    assert bledy


def test_plik_uczestnikow_xlsx_z_polskimi_naglowkami():
    plik = zbuduj_xlsx([
        {"Imię i nazwisko": "Łucja Żak", "E-mail": "lucja@firma.pl"},
        {"Imię i nazwisko": "Bez Maila", "E-mail": None},
    ])
    uczestnicy, bledy = db.wczytaj_plik_uczestnikow(plik.getvalue(), "lista.xlsx")
    assert bledy == []
    assert uczestnicy == [("Łucja Żak", "lucja@firma.pl"), ("Bez Maila", None)]


@pytest.mark.parametrize("kodowanie,separator", [("cp1250", ";"), ("utf-8-sig", ","), ("utf-8", "\t")])
def test_plik_uczestnikow_csv_w_roznych_kodowaniach(kodowanie, separator):
    dane = f"imie{separator}email\nŁucja Żak{separator}lucja@firma.pl\n".encode(kodowanie)
    uczestnicy, bledy = db.wczytaj_plik_uczestnikow(dane, "lista.csv")
    assert bledy == []
    assert uczestnicy == [("Łucja Żak", "lucja@firma.pl")]


def test_plik_uczestnikow_bez_naglowkow():
    plik = zbuduj_xlsx([{"cos": "Jan", "innego": "jan@firma.pl"}])
    uczestnicy, bledy = db.wczytaj_plik_uczestnikow(plik.getvalue(), "lista.xlsx")
    assert uczestnicy == []
    assert "brakuje nagłówków" in bledy[0]


def test_generowanie_tokenow_z_pliku(klient, haslo_admina, test_z_pytaniami):
    zaloguj_admina(klient, haslo_admina)
    plik = zbuduj_xlsx([{"imie": "Jan Kowalski", "email": "jan@firma.pl"}, {"imie": "Anna Nowak", "email": None}])
    odpowiedz = klient.post(
        f"/admin/testy/{test_z_pytaniami}/tokeny",
        data={"plik_uczestnikow": (plik, "lista.xlsx"), "dlugosc": "6", "csrf_token": pobierz_csrf_token(klient)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert odpowiedz.status_code == 200
    osoby = sorted((t["imie"], t["email"]) for t in _tokeny(test_z_pytaniami).values())
    assert osoby == [("Anna Nowak", None), ("Jan Kowalski", "jan@firma.pl")]


def test_bledna_lista_nie_generuje_tokenow_i_zostawia_tekst(klient, haslo_admina, test_z_pytaniami):
    zaloguj_admina(klient, haslo_admina)
    lista = "Jan;jan@firma.pl\nZły;zly@@adres"
    odpowiedz = klient.post(
        f"/admin/testy/{test_z_pytaniami}/tokeny",
        data={"lista_uczestnikow": lista, "dlugosc": "6", "csrf_token": pobierz_csrf_token(klient)},
    )
    tekst = odpowiedz.data.decode()
    assert odpowiedz.status_code == 200
    assert "linia 2" in tekst
    assert "Zły;zly@@adres" in tekst
    assert _tokeny(test_z_pytaniami) == {}


def test_migracja_pola_przypisany(tmp_path, monkeypatch):
    """Stara baza z kolumną `przypisany` — adres trafia do email, reszta do imie;
    druga inicjalizacja niczego nie psuje."""
    plik = tmp_path / "stara.db"
    conn = sqlite3.connect(plik)
    conn.execute(
        "CREATE TABLE tokeny (token TEXT PRIMARY KEY, test_id INTEGER NOT NULL, wykorzystany INTEGER NOT NULL DEFAULT 0, "
        "data_utworzenia TEXT NOT NULL, data_wykorzystania TEXT, przypisany TEXT)"
    )
    conn.executemany(
        "INSERT INTO tokeny (token, test_id, data_utworzenia, przypisany) VALUES (?, 1, '2026-01-01 10:00:00', ?)",
        [("T1", "Jan Kowalski"), ("T2", "Anna Nowak <anna@firma.pl>"), ("T3", "piotr@firma.pl"), ("T4", None)],
    )
    conn.commit()
    conn.close()

    monkeypatch.setattr(db, "DB_PLIK", str(plik))
    db.inicjalizuj()
    db.inicjalizuj()
    with db.baza() as conn:
        wiersze = {w["token"]: (w["imie"], w["email"]) for w in conn.execute("SELECT * FROM tokeny")}
    assert wiersze == {
        "T1": ("Jan Kowalski", None),
        "T2": ("Anna Nowak", "anna@firma.pl"),
        "T3": (None, "piotr@firma.pl"),
        "T4": (None, None),
    }


def test_przypisanie_zmiana_adresu_zeruje_status_maila(klient, haslo_admina, test_z_pytaniami):
    (wpis,) = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Jan", "jan@firma.pl")])
    token = wpis["token"]
    db.zapisz_wynik_wysylki(token, False)
    zaloguj_admina(klient, haslo_admina)

    def przypisz(imie, email):
        return klient.post(
            f"/admin/tokeny/{token}/przypisz",
            data={"imie": imie, "email": email, "test_id": str(test_z_pytaniami), "csrf_token": pobierz_csrf_token(klient)},
            follow_redirects=True,
        )

    przypisz("Jan Kowalski", "jan@firma.pl")
    assert _tokeny(test_z_pytaniami)[token]["mail_status"] == "ok"

    przypisz("Jan Kowalski", "jan.kowalski@firma.pl")
    stan = _tokeny(test_z_pytaniami)[token]
    assert (stan["email"], stan["mail_status"], stan["mail_wyslany_at"]) == ("jan.kowalski@firma.pl", None, None)

    odpowiedz = przypisz("Jan Kowalski", "to-nie-adres")
    assert "Nie zapisano przypisania" in odpowiedz.data.decode()
    assert _tokeny(test_z_pytaniami)[token]["email"] == "jan.kowalski@firma.pl"


# --- F3/E6: wysyłka -------------------------------------------------------------------

def test_wysylka_zaproszen(klient, haslo_admina, test_z_pytaniami, smtp):
    db.wygeneruj_tokeny(test_z_pytaniami, 3, dlugosc=6, uczestnicy=[
        ("Jan Kowalski", "jan@firma.pl"), ("Anna Nowak", "anna@firma.pl"), ("Bez Maila", None),
    ])
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = _wyslij(klient, test_z_pytaniami, "zaproszenia")
    tekst = odpowiedz.data.decode()
    assert "Wysyłka zakończona" in tekst
    assert "Wysłane: <strong>2</strong>" in tekst

    assert sorted(w["To"] for w in smtp.wyslane) == ["anna@firma.pl", "jan@firma.pl"]
    tokeny = _tokeny(test_z_pytaniami)
    for wiadomosc in smtp.wyslane:
        (token,) = [t for t, w in tokeny.items() if w["email"] == wiadomosc["To"]]
        tresc = wiadomosc.get_content()
        assert f"http://localhost:5555/?token={token}" in tresc
        assert "Powodzenia!" in tresc
        assert wiadomosc["Subject"] == "Zaproszenie do testu: Test testowy"
        assert "Test wiedzy" in wiadomosc["From"] and "nadawca@example.com" in wiadomosc["From"]
        assert tokeny[token]["mail_status"] == "ok" and tokeny[token]["mail_wyslany_at"]

    # Drugi raz zaproszenia nie idą do tych samych osób.
    odpowiedz = _wyslij(klient, test_z_pytaniami, "zaproszenia")
    assert "Brak odbiorców dla tej wysyłki" in odpowiedz.data.decode()
    assert len(smtp.wyslane) == 2


def test_przypomnienia_tylko_do_nieukonczonych(klient, haslo_admina, test_z_pytaniami, smtp):
    wpisy = db.wygeneruj_tokeny(test_z_pytaniami, 2, dlugosc=6, uczestnicy=[
        ("Jan", "jan@firma.pl"), ("Anna", "anna@firma.pl"),
    ])
    przejdz_caly_test(klient, wpisy[0]["token"])
    zaloguj_admina(klient, haslo_admina)
    _wyslij(klient, test_z_pytaniami, "przypomnienia")

    assert [w["To"] for w in smtp.wyslane] == ["anna@firma.pl"]
    assert smtp.wyslane[0]["Subject"] == "Przypomnienie: test Test testowy"
    tokeny = _tokeny(test_z_pytaniami)
    assert tokeny[wpisy[1]["token"]]["przypomnienie_wyslane_at"]
    assert tokeny[wpisy[1]["token"]]["mail_wyslany_at"] is None


def test_wysylka_do_zaznaczonych_ponownie(klient, haslo_admina, test_z_pytaniami, smtp):
    wpisy = db.wygeneruj_tokeny(test_z_pytaniami, 2, dlugosc=6, uczestnicy=[
        ("Jan", "jan@firma.pl"), ("Anna", "anna@firma.pl"),
    ])
    db.zapisz_wynik_wysylki(wpisy[0]["token"], False)
    zaloguj_admina(klient, haslo_admina)
    _wyslij(klient, test_z_pytaniami, "zaznaczone", [wpisy[0]["token"]])
    assert [w["To"] for w in smtp.wyslane] == ["jan@firma.pl"]

    odpowiedz = _wyslij(klient, test_z_pytaniami, "zaznaczone")
    assert "Zaznacz tokeny z adresem e-mail" in odpowiedz.data.decode()


def test_odrzucony_adres_dostaje_status_bledu(klient, haslo_admina, test_z_pytaniami, smtp):
    wpisy = db.wygeneruj_tokeny(test_z_pytaniami, 2, dlugosc=6, uczestnicy=[
        ("Jan", "nieistnieje@firma.pl"), ("Anna", "anna@firma.pl"),
    ])
    smtp.odrzucane_adresy.add("nieistnieje@firma.pl")
    zaloguj_admina(klient, haslo_admina)
    tekst = _wyslij(klient, test_z_pytaniami, "zaproszenia").data.decode()
    assert "Błędy: <strong>1</strong>" in tekst
    assert "550" in tekst

    tokeny = _tokeny(test_z_pytaniami)
    zly = tokeny[wpisy[0]["token"]]
    assert zly["mail_status"] == "blad" and "550" in zly["mail_blad"] and zly["mail_wyslany_at"] is None
    assert tokeny[wpisy[1]["token"]]["mail_status"] == "ok"
    assert [w["To"] for w in smtp.wyslane] == ["anna@firma.pl"]

    # Na liście tokenów widać błąd, a kolejne „zaproszenia” spróbują ponownie.
    lista = klient.get(f"/admin/testy/{test_z_pytaniami}/tokeny").data.decode()
    assert "błąd" in lista and "550" in lista
    assert [o["email"] for o in db.tokeny_do_wysylki(test_z_pytaniami, "zaproszenia")] == ["nieistnieje@firma.pl"]


def test_zle_haslo_smtp_nie_ujawnia_hasla(klient, haslo_admina, test_z_pytaniami, smtp):
    db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Jan", "jan@firma.pl")])
    smtp.zle_haslo = True
    zaloguj_admina(klient, haslo_admina)
    tekst = _wyslij(klient, test_z_pytaniami, "zaproszenia").data.decode()
    assert "hasła do aplikacji" in tekst
    assert HASLO_SMTP not in tekst
    (stan,) = _tokeny(test_z_pytaniami).values()
    assert stan["mail_status"] == "blad" and HASLO_SMTP not in stan["mail_blad"]


def test_wysylka_bez_config_json(klient, haslo_admina, test_z_pytaniami, bez_konfiguracji):
    db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Jan", "jan@firma.pl")])
    zaloguj_admina(klient, haslo_admina)
    tekst = _wyslij(klient, test_z_pytaniami, "zaproszenia").data.decode()
    assert "Brak pliku config.json" in tekst
    (stan,) = _tokeny(test_z_pytaniami).values()
    assert stan["mail_status"] is None

    strona = klient.get(f"/admin/testy/{test_z_pytaniami}/mail").data.decode()
    assert "Wysyłka z aplikacji jest wyłączona" in strona


def test_nieznana_wysylka_wraca_do_tokenow(klient, haslo_admina, test_z_pytaniami):
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get(f"/admin/wysylki/nieznana?test_id={test_z_pytaniami}")
    assert odpowiedz.status_code == 302
    assert f"/admin/testy/{test_z_pytaniami}/tokeny" in odpowiedz.headers["Location"]


# --- Szablony maili -------------------------------------------------------------------

def _zapisz_szablony(klient, test_id, **pola):
    dane = {"mail_temat": "", "mail_tresc": "", "przypomnienie_temat": "", "przypomnienie_tresc": ""}
    dane.update(pola)
    dane["csrf_token"] = pobierz_csrf_token(klient)
    return klient.post(f"/admin/testy/{test_id}/mail", data=dane, follow_redirects=True)


def test_wlasny_szablon_jest_uzywany_przy_wysylce(klient, haslo_admina, test_z_pytaniami, smtp):
    db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Jan Kowalski", "jan@firma.pl")])
    zaloguj_admina(klient, haslo_admina)
    klient.get(f"/admin/testy/{test_z_pytaniami}/mail")
    odpowiedz = _zapisz_szablony(
        klient, test_z_pytaniami, mail_temat="Cześć {imie}", mail_tresc="Twój token: {token}\n{link}"
    )
    assert "Zapisano szablony maili" in odpowiedz.data.decode()
    assert "Cześć Jan Kowalski" in odpowiedz.data.decode()  # podgląd na pierwszym odbiorcy

    _wyslij(klient, test_z_pytaniami, "zaproszenia")
    (wiadomosc,) = smtp.wyslane
    (token,) = _tokeny(test_z_pytaniami)
    assert wiadomosc["Subject"] == "Cześć Jan Kowalski"
    assert wiadomosc.get_content().strip() == f"Twój token: {token}\nhttp://localhost:5555/?token={token}"


def test_pusty_albo_domyslny_szablon_zapisuje_null(klient, haslo_admina, test_z_pytaniami):
    zaloguj_admina(klient, haslo_admina)
    klient.get(f"/admin/testy/{test_z_pytaniami}/mail")
    _zapisz_szablony(klient, test_z_pytaniami, mail_temat=poczta.DOMYSLNE_SZABLONY["zaproszenie"]["temat"])
    with db.baza() as conn:
        test = conn.execute("SELECT * FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()
    assert [test[k] for k in ("mail_temat", "mail_tresc", "przypomnienie_temat", "przypomnienie_tresc")] == [None] * 4


def test_nieznany_znacznik_w_szablonie_jest_odrzucany(klient, haslo_admina, test_z_pytaniami):
    zaloguj_admina(klient, haslo_admina)
    klient.get(f"/admin/testy/{test_z_pytaniami}/mail")
    odpowiedz = _zapisz_szablony(klient, test_z_pytaniami, mail_tresc="Witaj {imię} {nazwisko}")
    tekst = odpowiedz.data.decode()
    assert "nieznane znaczniki" in tekst and "{nazwisko}" in tekst
    assert "Witaj {imię} {nazwisko}" in tekst
    with db.baza() as conn:
        assert conn.execute("SELECT mail_tresc FROM testy WHERE id = ?", (test_z_pytaniami,)).fetchone()[0] is None


# --- Ustawienia poczty ----------------------------------------------------------------

def test_strona_poczty_nie_pokazuje_hasla(klient, haslo_admina, smtp):
    zaloguj_admina(klient, haslo_admina)
    tekst = klient.get("/admin/poczta").data.decode()
    assert "smtp.example.com:587" in tekst
    assert "ustawione" in tekst
    assert HASLO_SMTP not in tekst
    assert "localhost" in tekst and "działa tylko na komputerze" in tekst


def test_zapis_ustawien_poczty_normalizuje_adres(klient, haslo_admina):
    zaloguj_admina(klient, haslo_admina)
    klient.get("/admin/poczta")
    odpowiedz = klient.post("/admin/poczta", data={
        "nazwa_nadawcy": "  Dział   Szkoleń ", "adres_aplikacji": "192.168.1.10:5555/",
        "csrf_token": pobierz_csrf_token(klient),
    }, follow_redirects=True)
    assert "Zapisano ustawienia poczty" in odpowiedz.data.decode()
    assert db.pobierz_ustawienia_aplikacji() == {
        "nazwa_nadawcy": "Dział Szkoleń", "adres_aplikacji": "http://192.168.1.10:5555",
    }
    assert "działa tylko na komputerze" not in odpowiedz.data.decode()

    odpowiedz = klient.post("/admin/poczta", data={
        "nazwa_nadawcy": "X", "adres_aplikacji": "http://zły adres",
        "csrf_token": pobierz_csrf_token(klient),
    }, follow_redirects=True)
    assert "Nieprawidłowy adres aplikacji" in odpowiedz.data.decode()
    assert db.pobierz_ustawienia_aplikacji()["adres_aplikacji"] == "http://192.168.1.10:5555"


def test_mail_testowy(klient, haslo_admina, smtp):
    zaloguj_admina(klient, haslo_admina)
    klient.get("/admin/poczta")
    odpowiedz = klient.post("/admin/poczta/test", data={
        "adres": "ja@firma.pl", "csrf_token": pobierz_csrf_token(klient),
    }, follow_redirects=True)
    assert "Wysłano mail testowy na ja@firma.pl" in odpowiedz.data.decode()
    assert [w["To"] for w in smtp.wyslane] == ["ja@firma.pl"]

    odpowiedz = klient.post("/admin/poczta/test", data={
        "adres": "zly", "csrf_token": pobierz_csrf_token(klient),
    }, follow_redirects=True)
    assert "Nie wysłano maila testowego" in odpowiedz.data.decode()
    assert len(smtp.wyslane) == 1


def test_mail_testowy_zle_haslo(klient, haslo_admina, smtp):
    smtp.zle_haslo = True
    zaloguj_admina(klient, haslo_admina)
    klient.get("/admin/poczta")
    tekst = klient.post("/admin/poczta/test", data={
        "adres": "ja@firma.pl", "csrf_token": pobierz_csrf_token(klient),
    }, follow_redirects=True).data.decode()
    assert "hasła do aplikacji" in tekst
    assert HASLO_SMTP not in tekst


@pytest.mark.parametrize("sciezka", ["/admin/poczta", "/admin/poczta/test", "/admin/testy/{id}/mail", "/admin/testy/{id}/wyslij"])
def test_nowe_formularze_wymagaja_csrf(klient, haslo_admina, test_z_pytaniami, sciezka):
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.post(sciezka.format(id=test_z_pytaniami), data={"rodzaj": "zaproszenia"})
    assert odpowiedz.status_code == 400


def test_strony_maili_wymagaja_logowania(klient, test_z_pytaniami):
    for sciezka in ("/admin/poczta", f"/admin/testy/{test_z_pytaniami}/mail", f"/admin/testy/{test_z_pytaniami}/korespondencja.xlsx"):
        assert klient.get(sciezka).status_code == 302


# --- Token z linku, korespondencja seryjna, wynik jednej osoby ------------------------

def test_token_z_linku_jest_wpisany_w_formularz(klient):
    assert 'value="ABC234"' in klient.get("/?token=abc234").data.decode()
    tekst = klient.get('/?token="><script>alert(1)</script>').data.decode()
    assert "<script>alert" not in tekst
    assert 'value="SCRIPTALERT1SCRIPT"' in tekst


def test_eksport_do_korespondencji_seryjnej(klient, haslo_admina, test_z_pytaniami):
    wpisy = db.wygeneruj_tokeny(test_z_pytaniami, 3, dlugosc=6, uczestnicy=[
        ("Jan", "jan@firma.pl"), ("Anna", "anna@firma.pl"), ("Bez Maila", None),
    ])
    przejdz_caly_test(klient, wpisy[1]["token"])
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get(f"/admin/testy/{test_z_pytaniami}/korespondencja.xlsx")
    assert odpowiedz.status_code == 200
    wiersze = list(_xlsx(odpowiedz)["Korespondencja"].iter_rows(values_only=True))
    assert wiersze[0] == ("imie", "email", "token", "link", "nazwa_testu", "dostepny_do", "limit_czasu")
    token = wpisy[0]["token"]
    assert wiersze[1:] == [("Jan", "jan@firma.pl", token, f"http://localhost:5555/?token={token}",
                            "Test testowy", "bez terminu", "bez limitu czasu")]


def test_eksport_wyniku_jednej_osoby(klient, haslo_admina, test_z_pytaniami):
    (wpis,) = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6, uczestnicy=[("Łucja Żak", "lucja@firma.pl")])
    przejdz_caly_test(klient, wpis["token"], "B")
    zaloguj_admina(klient, haslo_admina)

    strona = klient.get(f"/admin/tokeny/{wpis['token']}").data.decode()
    assert "3 / 3 (100.0%)" in strona
    assert f"/admin/tokeny/{wpis['token']}/wynik.xlsx" in strona

    odpowiedz = klient.get(f"/admin/tokeny/{wpis['token']}/wynik.xlsx")
    assert odpowiedz.status_code == 200
    assert f"wynik_{wpis['token']}.xlsx" in odpowiedz.headers["Content-Disposition"]
    wiersze = list(_xlsx(odpowiedz)["Wynik"].iter_rows(values_only=True))
    naglowek = {w[0]: w[1] for w in wiersze[:11]}
    assert naglowek["Imię"] == "Łucja Żak"
    assert naglowek["E-mail"] == "lucja@firma.pl"
    assert naglowek["Status"] == "ukończony"
    assert naglowek["Wynik"] == "3 / 3"
    assert naglowek["Procent"] == 1.0
    assert naglowek["Zaliczenie"] == "zdał (próg 80%)"
    assert isinstance(naglowek["Rozpoczęto"], datetime)
    tabela = wiersze[wiersze.index(("Nr", "Pytanie", "Odpowiedź uczestnika", "Poprawna odpowiedź", "Wynik", "Wyjaśnienie")) + 1:]
    assert [w[0] for w in tabela] == [1, 2, 3]
    assert all(w[4] == "poprawna" for w in tabela)


def test_wynik_osoby_przy_czasie_minal_liczy_wszystkie_pytania(klient, haslo_admina, test_z_pytaniami):
    """Mianownik to liczba pytań podejścia, a nie liczba udzielonych odpowiedzi."""
    db.zapisz_ustawienia_testu(test_z_pytaniami, 1, "po_zamknieciu", None, "od_razu", None, None)
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    odpowiedz = klient.post("/", data={"token": token}, follow_redirects=True)
    odpowiedz = klient.post("/test", data={"start": "1"}, follow_redirects=True)
    pytanie_id = re.search(rb'name="pytanie_id" value="(\d+)"', odpowiedz.data).group(1).decode()
    klient.post("/test", data={"pytanie_id": pytanie_id, "odpowiedz": "B"})
    stara = (datetime.now() - timedelta(minutes=5)).strftime(db.FORMAT_DATY)
    with db.baza() as conn:
        conn.execute("UPDATE podejscia SET data_rozpoczecia = ? WHERE token = ?", (stara, token))

    zaloguj_admina(klient, haslo_admina)
    strona = klient.get(f"/admin/tokeny/{token}").data.decode()
    assert "1 / 3 (33.3%)" in strona
    assert "czas minął" in strona
    assert "Bez odpowiedzi (np. minął czas): 2" in strona

    wiersze = list(_xlsx(klient.get(f"/admin/tokeny/{token}/wynik.xlsx"))["Wynik"].iter_rows(values_only=True))
    naglowek = {w[0]: w[1] for w in wiersze[:11]}
    assert naglowek["Wynik"] == "1 / 3"
    assert naglowek["Zaliczenie"] == "nie zdał (próg 80%)"
    tabela = wiersze[wiersze.index(("Nr", "Pytanie", "Odpowiedź uczestnika", "Poprawna odpowiedź", "Wynik", "Wyjaśnienie")) + 1:]
    assert len(tabela) == 3
    assert sum(1 for w in tabela if w[2] == "brak odpowiedzi") == 2


def test_wynik_osoby_nieukonczony_nie_ma_eksportu(klient, haslo_admina, test_z_pytaniami):
    token = db.wygeneruj_tokeny(test_z_pytaniami, 1, dlugosc=6)[0]["token"]
    zaloguj_admina(klient, haslo_admina)
    odpowiedz = klient.get(f"/admin/tokeny/{token}/wynik.xlsx")
    assert odpowiedz.status_code == 302
    assert f"/admin/testy/{test_z_pytaniami}/tokeny" in odpowiedz.headers["Location"]
